import os
import shutil
import logging
import urllib.parse
import subprocess
from typing import Dict, Any, Optional, List
from pathlib import Path

from ..database import DATA_DIR
from ..models import File

logger = logging.getLogger("memora.preview_service")

PREVIEW_CACHE_DIR = os.path.join(DATA_DIR, "preview_cache")
os.makedirs(PREVIEW_CACHE_DIR, exist_ok=True)

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".svg", ".tiff"}
PRESENTATION_EXTENSIONS = {".pptx", ".ppt"}
DOCUMENT_EXTENSIONS = {".docx", ".doc"}
PDF_EXTENSIONS = {".pdf"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".html", ".htm",
    ".js", ".jsx", ".ts", ".tsx", ".py", ".css", ".sql", ".sh", ".bat", ".ps1", ".log"
}


def get_soffice_path() -> Optional[str]:
    """Finds LibreOffice soffice executable on Windows or POSIX."""
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    found = shutil.which("soffice")
    if found and os.path.exists(found):
        return found
    return None


class PreviewService:
    def __init__(self):
        self.soffice_path = get_soffice_path()
        if self.soffice_path:
            logger.info(f"LibreOffice detected at '{self.soffice_path}' for document & presentation conversion.")
        else:
            logger.warning("LibreOffice not detected. Falling back to native python parsers for presentations and documents.")

    def convert_to_pdf(self, file_path: str, cache_key: str) -> Optional[str]:
        """Converts an office document to PDF via headless LibreOffice and caches it."""
        if not self.soffice_path or not os.path.exists(file_path):
            return None

        cached_pdf_name = f"{cache_key}.pdf"
        cached_pdf_path = os.path.join(PREVIEW_CACHE_DIR, cached_pdf_name)

        if os.path.exists(cached_pdf_path) and os.path.getsize(cached_pdf_path) > 0:
            return cached_pdf_path

        try:
            cmd = [
                self.soffice_path,
                "--headless",
                "--convert-to", "pdf",
                os.path.abspath(file_path),
                "--outdir", PREVIEW_CACHE_DIR
            ]
            logger.info(f"Running document conversion: {file_path} -> PDF")
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=35)
            if res.returncode == 0:
                # LibreOffice outputs with same stem name as original file
                orig_stem = Path(file_path).stem
                default_output = os.path.join(PREVIEW_CACHE_DIR, f"{orig_stem}.pdf")
                if os.path.exists(default_output):
                    if default_output != cached_pdf_path:
                        if os.path.exists(cached_pdf_path):
                            try:
                                os.remove(cached_pdf_path)
                            except Exception:
                                pass
                        shutil.move(default_output, cached_pdf_path)
                    logger.info(f"Successfully converted to PDF preview: {cached_pdf_path}")
                    return cached_pdf_path
            else:
                logger.warning(f"LibreOffice conversion failed with exit code {res.returncode}: {res.stderr}")
        except Exception as err:
            logger.error(f"Error converting document to PDF: {err}", exc_info=True)

        return None

    def get_file_preview_data(self, file_rec: File) -> Dict[str, Any]:
        """
        Inspects and returns visual preview data for the file.
        Supports PowerPoint (visual slides via PDF or structured slide deck),
        PDF, Images, Word documents, text/code, and clean unsupported fallbacks.
        """
        if not file_rec or not file_rec.path or not os.path.exists(file_rec.path):
            return {
                "format": "missing",
                "message": "File does not exist on local disk.",
                "filename": file_rec.name if file_rec else "Unknown",
                "path": file_rec.path if file_rec else ""
            }

        file_path = file_rec.path
        ext = os.path.splitext(file_rec.name)[1].lower()
        mtime = int(file_rec.modified_at.timestamp()) if file_rec.modified_at else int(os.path.getmtime(file_path))
        cache_key = f"preview_{file_rec.id}_{mtime}"

        # ---------------------------------------------------------------------
        # 1. Images (JPG, PNG, WEBP, GIF, SVG, BMP)
        # ---------------------------------------------------------------------
        if ext in IMAGE_EXTENSIONS:
            dims = "Auto"
            try:
                from PIL import Image
                with Image.open(file_path) as img:
                    dims = f"{img.width} × {img.height} px"
            except Exception:
                pass

            return {
                "format": "image",
                "url": f"/api/media/{file_rec.id}/preview",
                "dimensions": dims,
                "filename": file_rec.name,
                "filesize": file_rec.size,
                "path": file_path
            }

        # ---------------------------------------------------------------------
        # 2. PDF Documents
        # ---------------------------------------------------------------------
        if ext in PDF_EXTENSIONS:
            total_pages = 1
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                total_pages = len(reader.pages)
            except Exception as e:
                logger.debug(f"Failed to count PDF pages via pypdf: {e}")

            return {
                "format": "pdf",
                "url": f"/api/pdf/preview-file?file_path={urllib.parse.quote(file_path)}",
                "total_pages": total_pages,
                "filename": file_rec.name,
                "filesize": file_rec.size,
                "path": file_path
            }

        # ---------------------------------------------------------------------
        # 3. PowerPoint Presentations (PPTX, PPT)
        # ---------------------------------------------------------------------
        if ext in PRESENTATION_EXTENSIONS:
            # Attempt 1: High-fidelity visual PDF slide conversion via LibreOffice
            pdf_path = self.convert_to_pdf(file_path, cache_key)
            if pdf_path and os.path.exists(pdf_path):
                total_slides = 1
                try:
                    import pypdf
                    reader = pypdf.PdfReader(pdf_path)
                    total_slides = len(reader.pages)
                except Exception:
                    pass

                # Extract slide titles with python-pptx if available
                slide_titles = []
                if ext == ".pptx":
                    try:
                        from pptx import Presentation
                        prs = Presentation(file_path)
                        for s in prs.slides:
                            t = ""
                            if s.shapes.title and hasattr(s.shapes.title, "text"):
                                t = s.shapes.title.text.strip()
                            slide_titles.append(t or f"Slide {len(slide_titles) + 1}")
                    except Exception:
                        pass

                return {
                    "format": "presentation_pdf",
                    "pdf_url": f"/api/pdf/preview-file?file_path={urllib.parse.quote(pdf_path)}",
                    "total_slides": total_slides,
                    "slide_titles": slide_titles,
                    "filename": file_rec.name,
                    "filesize": file_rec.size,
                    "path": file_path
                }

            # Attempt 2: Structured in-app slide deck parser via python-pptx
            if ext == ".pptx":
                try:
                    from pptx import Presentation
                    prs = Presentation(file_path)
                    slides_data = []

                    for idx, slide in enumerate(prs.slides, start=1):
                        title = ""
                        bullets = []
                        tables = []

                        if slide.shapes.title and hasattr(slide.shapes.title, "text"):
                            title = slide.shapes.title.text.strip()

                        for shape in slide.shapes:
                            # Normal text boxes and shapes
                            if shape != slide.shapes.title and hasattr(shape, "text") and shape.text:
                                text = shape.text.strip()
                                if text:
                                    bullets.append(text)

                            # Tables
                            if getattr(shape, "has_table", False):
                                table_rows = []
                                for row in shape.table.rows:
                                    row_vals = [cell.text.strip() for cell in row.cells]
                                    if any(row_vals):
                                        table_rows.append(row_vals)
                                if table_rows:
                                    tables.append(table_rows)

                        slides_data.append({
                            "slide_number": idx,
                            "title": title or f"Slide {idx}",
                            "bullets": bullets,
                            "tables": tables
                        })

                    return {
                        "format": "presentation_slides",
                        "total_slides": len(slides_data),
                        "slides": slides_data,
                        "filename": file_rec.name,
                        "filesize": file_rec.size,
                        "path": file_path
                    }
                except Exception as p_err:
                    logger.error(f"Failed to parse PPTX with python-pptx: {p_err}", exc_info=True)

        # ---------------------------------------------------------------------
        # 4. Word Documents (DOCX, DOC)
        # ---------------------------------------------------------------------
        if ext in DOCUMENT_EXTENSIONS:
            # Attempt 1: High-fidelity visual PDF document conversion
            pdf_path = self.convert_to_pdf(file_path, cache_key)
            if pdf_path and os.path.exists(pdf_path):
                total_pages = 1
                try:
                    import pypdf
                    reader = pypdf.PdfReader(pdf_path)
                    total_pages = len(reader.pages)
                except Exception:
                    pass

                return {
                    "format": "document_pdf",
                    "pdf_url": f"/api/pdf/preview-file?file_path={urllib.parse.quote(pdf_path)}",
                    "total_pages": total_pages,
                    "filename": file_rec.name,
                    "filesize": file_rec.size,
                    "path": file_path
                }

            # Attempt 2: Structured DOCX parser via python-docx
            if ext == ".docx":
                try:
                    import docx
                    doc = docx.Document(file_path)
                    paragraphs = []
                    for p in doc.paragraphs:
                        text = p.text.strip()
                        if text:
                            style_name = p.style.name.lower() if p.style else ""
                            is_heading = "heading" in style_name or "title" in style_name
                            paragraphs.append({
                                "text": text,
                                "is_heading": is_heading,
                                "style": style_name
                            })

                    tables_data = []
                    for t in doc.tables:
                        t_rows = []
                        for row in t.rows:
                            r_vals = [cell.text.strip() for cell in row.cells]
                            if any(r_vals):
                                t_rows.append(r_vals)
                        if t_rows:
                            tables_data.append(t_rows)

                    return {
                        "format": "document_rich",
                        "paragraphs": paragraphs,
                        "tables": tables_data,
                        "filename": file_rec.name,
                        "filesize": file_rec.size,
                        "path": file_path
                    }
                except Exception as d_err:
                    logger.error(f"Failed to parse DOCX with python-docx: {d_err}", exc_info=True)

        # ---------------------------------------------------------------------
        # 5. Text / Code / Markdown / CSV
        # ---------------------------------------------------------------------
        if ext in TEXT_EXTENSIONS:
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    lines = [next(f) for _ in range(1200)]
                    raw_text = "".join(lines)
                return {
                    "format": "text",
                    "content": raw_text,
                    "line_count": len(lines),
                    "filename": file_rec.name,
                    "filesize": file_rec.size,
                    "path": file_path
                }
            except Exception as t_err:
                logger.error(f"Failed to read text file: {t_err}")

        # ---------------------------------------------------------------------
        # 6. Fallback for other formats
        # ---------------------------------------------------------------------
        return {
            "format": "unsupported",
            "extension": ext.upper().replace(".", "") if ext else "FILE",
            "filename": file_rec.name,
            "filesize": file_rec.size,
            "path": file_path,
            "message": f"Visual in-app preview is not supported for {ext.upper()} format."
        }


preview_service = PreviewService()
