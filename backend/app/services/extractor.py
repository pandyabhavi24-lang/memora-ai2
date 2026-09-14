import os
import logging
from typing import Tuple

logger = logging.getLogger("memora.extractor")

_easyocr_reader = None


def get_easyocr_reader():
    global _easyocr_reader

    if _easyocr_reader is None:
        try:
            import easyocr

            logger.info("Initializing EasyOCR reader (English)...")
            _easyocr_reader = easyocr.Reader(["en"], gpu=False)
            logger.info("EasyOCR reader initialized.")

        except Exception as e:
            logger.error(
                f"Failed to initialize EasyOCR reader: {e}",
                exc_info=True
            )
            _easyocr_reader = False

    return _easyocr_reader if _easyocr_reader is not False else None


class TextExtractor:
    """
    Extracts text from:

    - PDF
    - DOCX
    - PPTX
    - TXT / MD / CSV
    - Images using OCR

    Returns:
        (extracted_text, status)

    status:
        success
        empty
        failed
        skipped
    """

    @staticmethod
    def extract_from_pdf(file_path: str) -> Tuple[str, str]:
        try:
            import pypdf

            reader = pypdf.PdfReader(file_path)
            pages_text = []

            for i, page in enumerate(reader.pages):
                text = page.extract_text()

                if text and text.strip():
                    pages_text.append(text.strip())

            full_text = "\n\n".join(pages_text).strip()

            if full_text:
                return full_text, "success"

            # -------------------------------------------------------------
            # Scanned / Image-only PDF Fallback: Run existing EasyOCR
            # -------------------------------------------------------------
            logger.info(f"No embedded text found in PDF '{file_path}'. Processing scanned PDF via EasyOCR...")
            ocr_pages_text = []
            
            import io
            import tempfile
            from PIL import Image

            for idx, page in enumerate(reader.pages):
                page_ocr_text = []

                # Extract embedded images on this page
                if hasattr(page, "images") and page.images:
                    for img_count, img_obj in enumerate(page.images):
                        try:
                            with Image.open(io.BytesIO(img_obj.data)) as pil_img:
                                tmp_fd, tmp_path = tempfile.mkstemp(suffix=".png")
                                os.close(tmp_fd)
                                pil_img.save(tmp_path, format="PNG")

                                extracted_ocr, st = TextExtractor.extract_from_image(tmp_path)

                                if os.path.exists(tmp_path):
                                    try:
                                        os.remove(tmp_path)
                                    except Exception:
                                        pass

                                if extracted_ocr and extracted_ocr.strip():
                                    page_ocr_text.append(extracted_ocr.strip())
                        except Exception as img_err:
                            logger.debug(f"OCR image extraction failed for page {idx + 1} img {img_count + 1}: {img_err}")

                if page_ocr_text:
                    ocr_pages_text.append(" ".join(page_ocr_text))

            ocr_full_text = "\n\n".join(ocr_pages_text).strip()

            if not ocr_full_text:
                logger.info(f"No OCR text found in scanned PDF '{file_path}'.")
                return "", "empty"

            logger.info(f"OCR successfully extracted text from scanned PDF '{file_path}' ({len(ocr_full_text)} chars).")
            return ocr_full_text, "success"

        except Exception as e:
            logger.error(
                f"PDF extraction failed for '{file_path}': {e}"
            )
            return "", "failed"


    @staticmethod
    def extract_from_docx(file_path: str) -> Tuple[str, str]:
        try:
            import docx

            doc = docx.Document(file_path)

            paragraphs = [
                p.text.strip()
                for p in doc.paragraphs
                if p.text and p.text.strip()
            ]

            # Extract table text
            table_texts = []

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]
                    )

                    if row_text:
                        table_texts.append(row_text)

            full_text = "\n".join(
                paragraphs + table_texts
            ).strip()

            if not full_text:
                return "", "empty"

            return full_text, "success"

        except Exception as e:
            logger.error(
                f"DOCX extraction failed for '{file_path}': {e}"
            )
            return "", "failed"

    @staticmethod
    def extract_from_pptx(file_path: str) -> Tuple[str, str]:
        """
        Extract text from a PowerPoint .pptx file.

        Extracts:
        - slide text
        - text boxes
        - titles
        - tables
        - speaker notes where available

        Slide numbers are preserved so the extracted content
        remains understandable during semantic search.
        """

        try:
            from pptx import Presentation

            logger.info(
                f"Extracting PowerPoint text from '{file_path}'"
            )

            presentation = Presentation(file_path)

            slides_text = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1
            ):
                slide_parts = []

                # -------------------------------------------------
                # Extract normal slide shapes/text
                # -------------------------------------------------

                for shape in slide.shapes:

                    # Normal text boxes, titles, placeholders, etc.
                    if hasattr(shape, "text"):
                        text = shape.text.strip()

                        if text:
                            slide_parts.append(text)

                    # -------------------------------------------------
                    # Extract table contents
                    # -------------------------------------------------

                    if getattr(shape, "has_table", False):
                        table = shape.table

                        for row in table.rows:
                            row_text = " | ".join(
                                [
                                    cell.text.strip()
                                    for cell in row.cells
                                    if cell.text.strip()
                                ]
                            )

                            if row_text:
                                slide_parts.append(row_text)

                # -------------------------------------------------
                # Extract speaker notes
                # -------------------------------------------------

                try:
                    if slide.has_notes_slide:
                        notes_slide = slide.notes_slide

                        for shape in notes_slide.shapes:
                            if hasattr(shape, "text"):
                                notes_text = shape.text.strip()

                                if notes_text:
                                    slide_parts.append(
                                        notes_text
                                    )

                except Exception as notes_error:
                    # Notes extraction should never cause the
                    # entire PPTX extraction to fail.
                    logger.debug(
                        f"Could not extract notes from slide "
                        f"{slide_number}: {notes_error}"
                    )

                # -------------------------------------------------
                # Build slide content
                # -------------------------------------------------

                if slide_parts:
                    slide_text = (
                        f"[Slide {slide_number}]\n"
                        + "\n".join(slide_parts)
                    )

                    slides_text.append(slide_text)

            full_text = "\n\n".join(slides_text).strip()

            if not full_text:
                logger.info(
                    f"No extractable text found in PPTX "
                    f"'{file_path}'"
                )
                return "", "empty"

            logger.info(
                f"Successfully extracted PPTX '{file_path}' "
                f"({len(presentation.slides)} slides)"
            )

            return full_text, "success"

        except ImportError:
            logger.error(
                "python-pptx is not installed. "
                "Install it using: pip install python-pptx"
            )
            return "", "failed"

        except Exception as e:
            logger.error(
                f"PPTX extraction failed for '{file_path}': {e}",
                exc_info=True
            )
            return "", "failed"

    @staticmethod
    def extract_from_plain_text(file_path: str) -> Tuple[str, str]:
        encodings = [
            "utf-8",
            "utf-8-sig",
            "latin-1",
            "cp1252"
        ]

        for enc in encodings:
            try:
                with open(
                    file_path,
                    "r",
                    encoding=enc,
                    errors="replace"
                ) as f:

                    content = f.read().strip()

                    if content:
                        return content, "success"

                    return "", "empty"

            except Exception:
                continue

        return "", "failed"

    IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".tif", ".gif"}

    @staticmethod
    def extract_from_image(file_path: str) -> Tuple[str, str]:
        """
        Extract text from images using EasyOCR with robust format and path handling.
        Returns (extracted_text, status): 'success', 'empty', or 'failed'.
        """
        try:
            from PIL import Image
            import numpy as np

            reader = get_easyocr_reader()

            if not reader:
                logger.warning(
                    f"EasyOCR reader unavailable for image '{file_path}'"
                )
                return "", "failed"

            if not os.path.exists(file_path):
                logger.warning(f"Image file does not exist: '{file_path}'")
                return "", "failed"

            # Load image safely via PIL (handles JPG, PNG, WEBP, BMP, TIFF, GIF)
            try:
                with Image.open(file_path) as pil_img:
                    rgb_img = pil_img.convert("RGB")
                    img_array = np.array(rgb_img)
            except Exception as load_err:
                logger.debug(f"PIL Image.open failed for '{file_path}': {load_err}, attempting fallback")
                import cv2
                img_bytes = np.fromfile(file_path, dtype=np.uint8)
                img_array = cv2.imdecode(img_bytes, cv2.IMREAD_COLOR)

            if img_array is None or img_array.size == 0:
                logger.warning(f"Could not decode image pixel data for '{file_path}'")
                return "", "failed"

            # Run EasyOCR
            results = reader.readtext(img_array, detail=0)

            # Filter out whitespace and empty entries
            clean_results = [
                res.strip()
                for res in results
                if res and res.strip()
            ]

            extracted_str = " ".join(clean_results).strip()

            if not extracted_str:
                logger.debug(f"No text detected via OCR in image '{file_path}' (status: empty)")
                return "", "empty"

            logger.info(f"OCR successfully extracted {len(clean_results)} text segments from '{file_path}' ({len(extracted_str)} chars)")
            return extracted_str, "success"

        except Exception as e:
            logger.error(
                f"OCR extraction failed for '{file_path}': {e}",
                exc_info=True
            )
            return "", "failed"

    @classmethod
    def extract(
        cls,
        file_path: str,
        extension: str
    ) -> Tuple[str, str]:

        ext = extension.lower().strip()

        if not ext.startswith("."):
            ext = "." + ext

        if ext == ".pdf":
            return cls.extract_from_pdf(file_path)

        elif ext in [".docx", ".doc"]:
            return cls.extract_from_docx(file_path)

        elif ext == ".pptx":
            return cls.extract_from_pptx(file_path)

        elif ext in [".txt", ".md", ".csv", ".java", ".c", ".py", ".class", ".html", ".htm", ".js", ".ts", ".css", ".scss", ".json", ".xml", ".yml", ".yaml", ".sql"]:
            return cls.extract_from_plain_text(file_path)

        elif ext in cls.IMAGE_EXTENSIONS:
            return cls.extract_from_image(file_path)

        else:
            return "", "skipped"


text_extractor = TextExtractor()